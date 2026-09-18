import os
import tempfile
import pytest
from pptx import Presentation

try:
    import pypandoc
except ImportError:
    pypandoc = None

from opendocagent.converters import DocxConverter, PptxConverter, LatexConverter
from opendocagent.parser import MarkdownParser
from opendocagent.exceptions import OpenDocAgentError


def test_docx_converter():
    converter = DocxConverter()
    
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        out_path = tmp.name
        
    try:
        converter.convert({"metadata": {}, "tokens": []}, out_path)
        assert os.path.exists(out_path)
        assert os.path.getsize(out_path) > 0
    finally:
        if os.path.exists(out_path):
            os.remove(out_path)


def test_pptx_converter_basic():
    converter = PptxConverter()
    
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        out_path = tmp.name
        
    try:
        converter.convert({"metadata": {}, "tokens": []}, out_path)
        assert os.path.exists(out_path)
        assert os.path.getsize(out_path) > 0
    finally:
        if os.path.exists(out_path):
            os.remove(out_path)


def test_pptx_converter_rich_fidelity(tmp_path):
    parser = MarkdownParser()
    md = """---
title: "Quarterly Review"
format: "pptx"
style: "executive"
---

# Titan Enterprise Strategy
Next-Generation Agentic Document Pipelines

---

## Market Positioning
- Growing adoption in autonomous agents
  - High performance requirements
  - Enterprise data compliance
- Low barrier to entry

> Note: Highlight that our open-source agent pipeline runs without vendor lock-in.

---

## Architecture Pillars

Core Engine:
- Python >= 3.12
- In-memory streaming

::: col

Agentic Protocols:
- FastMCP tools
- Standard diagnostics
"""
    parsed = parser.parse(md)
    out_path = str(tmp_path / "deck.pptx")

    converter = PptxConverter()
    converter.convert(parsed, out_path)

    assert os.path.exists(out_path)
    assert os.path.getsize(out_path) > 0

    # Inspect generated presentation structures with python-pptx
    prs = Presentation(out_path)
    assert len(prs.slides) == 3

    # Slide 1: Title & Subtitle
    slide1 = prs.slides[0]
    assert slide1.shapes.title.text == "Titan Enterprise Strategy"
    if len(slide1.placeholders) > 1:
        assert "Next-Generation Agentic Document Pipelines" in slide1.placeholders[1].text

    # Slide 2: Heading, Bullet Hierarchy, Speaker Notes
    slide2 = prs.slides[1]
    assert slide2.shapes.title.text == "Market Positioning"
    # Check speaker notes extraction
    assert slide2.has_notes_slide is True
    notes_text = slide2.notes_slide.notes_text_frame.text
    assert "Highlight that our open-source agent pipeline" in notes_text
    assert not notes_text.startswith(">")

    # Check bullet levels
    body_tf = slide2.placeholders[1].text_frame
    paragraphs = [p for p in body_tf.paragraphs if p.text.strip()]
    assert any("Growing adoption" in p.text and p.level == 0 for p in paragraphs)
    assert any("High performance" in p.text and p.level == 1 for p in paragraphs)

    # Slide 3: Two-Column Structure
    slide3 = prs.slides[2]
    assert slide3.shapes.title.text == "Architecture Pillars"
    # Verify multiple shapes exist (left placeholder + right textbox)
    assert len(slide3.shapes) >= 2


def test_latex_converter():
    if pypandoc is None:
        pytest.skip("pypandoc not installed, skipping LaTeX test.")

    try:
        pypandoc.get_pandoc_version()
    except OSError:
        pytest.skip("Pandoc binary not found, skipping LaTeX integration test.")

    converter = LatexConverter()
    
    with tempfile.NamedTemporaryFile(suffix=".tex", delete=False) as tmp:
        out_path = tmp.name
        
    try:
        converter.convert({"metadata": {}, "tokens": []}, out_path)
        assert os.path.exists(out_path)
        assert os.path.getsize(out_path) > 0
    finally:
        if os.path.exists(out_path):
            try:
                os.remove(out_path)
            except Exception:
                pass


def test_latex_converter_missing_pypandoc(monkeypatch):
    import opendocagent.converters.latex as latex_mod
    monkeypatch.setattr(latex_mod, "pypandoc", None)

    conv = latex_mod.LatexConverter()
    with pytest.raises(OpenDocAgentError) as exc:
        conv.convert({"metadata": {}, "tokens": []}, "dummy.pdf")
    assert "pypandoc" in str(exc.value)
