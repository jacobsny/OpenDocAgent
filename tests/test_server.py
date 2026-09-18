"""Tests for OpenDocAgent FastMCP server tools, resources, and prompts."""

import json
from opendocagent.server import (
    lint_document,
    compile_document,
    get_style_guidelines,
    list_supported_formats,
    get_syntax_docs,
    get_architecture_docs,
    draft_document,
)


def test_mcp_lint_document_valid():
    md = """---
title: "Technical Specification"
author: "Lead Engineer"
format: "pdf"
style: "technical"
---
## Purpose
Define requirements.
## Scope
Core modules.
## Methodology
Step-by-step.
"""
    result = lint_document(md)
    assert result["valid"] is True
    assert result["errors"] == 0


def test_mcp_lint_document_invalid():
    md = """# No frontmatter here"""
    result = lint_document(md)
    assert result["valid"] is False
    assert result["errors"] > 0
    assert any(issue["code"] == "F001" for issue in result["issues"])


def test_mcp_compile_document(tmp_path):
    out_file = str(tmp_path / "output.docx")
    md = """---
title: "Title"
author: "Author"
format: "docx"
style: "executive"
---
# Heading
Body text.
"""
    res = compile_document(md, output_path=out_file, output_format="docx")
    assert res["success"] is True
    assert (tmp_path / "output.docx").exists()


def test_mcp_style_guidelines():
    guidance = get_style_guidelines("executive")
    assert "EXECUTIVE" in guidance
    assert "Pyramid Principle" in guidance


def test_mcp_supported_formats():
    formats = list_supported_formats()
    assert "docx" in formats
    assert "pdf" in formats
    assert "pptx" in formats


def test_mcp_resources():
    syntax = get_syntax_docs()
    assert len(syntax) > 0
    arch = get_architecture_docs()
    assert len(arch) > 0


def test_mcp_prompts():
    prompt = draft_document(style="executive", format="docx")
    assert "Target Format: docx" in prompt
    assert "EXECUTIVE" in prompt


def test_mcp_compile_document_pptx(tmp_path):
    from pptx import Presentation

    out_file = str(tmp_path / "deck.pptx")
    md = """---
title: "Agentic Presentation"
author: "FastMCP Agent"
format: "pptx"
style: "executive"
---

# Strategic Architecture
Next Gen AI Workflows

---

## Capabilities
- Multi-format generation
- FastMCP tool exposure

> Note: Highlight that PPTX generation supports full presenter notes.
"""
    res = compile_document(md, output_path=out_file, output_format="pptx")
    assert res["success"] is True
    assert (tmp_path / "deck.pptx").exists()

    prs = Presentation(out_file)
    assert len(prs.slides) == 2
    assert prs.slides[0].shapes.title.text == "Strategic Architecture"
    assert prs.slides[1].has_notes_slide is True
    assert "presenter notes" in prs.slides[1].notes_slide.notes_text_frame.text


def test_mcp_compile_document_auto_format(tmp_path):
    out_file = str(tmp_path / "auto_deck.pptx")
    md = """---
title: "Auto Format Deck"
format: "pptx"
style: "executive"
---

## Key Milestone
Content for auto-detected format.
"""
    # output_format defaults to "auto"
    res = compile_document(md, output_path=out_file)
    assert res["success"] is True
    assert (tmp_path / "auto_deck.pptx").exists()


def test_mcp_compile_document_template_override(tmp_path):
    out_file = str(tmp_path / "branded_deck.pptx")
    md = """---
title: "Branded Deck"
format: "pptx"
style: "executive"
---

## Executive Overview
Branded template test.
"""
    res = compile_document(
        md,
        output_path=out_file,
        output_format="pptx",
        template_override="executive_deck",
    )
    assert res["success"] is True
    assert (tmp_path / "branded_deck.pptx").exists()


def test_mcp_compile_document_missing_template(tmp_path):
    out_file = str(tmp_path / "fail.pptx")
    md = """---
title: "Fail Deck"
format: "pptx"
---
## Fail
"""
    res = compile_document(
        md,
        output_path=out_file,
        output_format="pptx",
        template_override="nonexistent_template_xyz",
    )
    assert res["success"] is False
    assert "not found" in res["error"]


def test_mcp_lint_document_malformed_json():
    result = lint_document("# Title\n", context_json="{broken_json: 123}")
    assert result["valid"] is False
    assert result["errors"] == 1
    assert result["issues"][0]["code"] == "E_JSON"


def test_mcp_compile_document_malformed_json(tmp_path):
    out_file = str(tmp_path / "fail.docx")
    res = compile_document(
        "# Title\n",
        output_path=out_file,
        output_format="docx",
        context_json="{bad_json}",
    )
    assert res["success"] is False
    assert res["error"] is not None

